#!/usr/bin/env python3
"""Build editable Chinese Word and matching PDF from the current manuscript.

Authoring only: no simulator imports or training. python-docx==1.2.0 and the
existing lxml/Pillow are required; PDF rendering uses local LibreOffice.
"""
from pathlib import Path
import argparse
import hashlib
import json
import os
import re
import subprocess
import tempfile

from docx import Document
from docx.enum.section import WD_SECTION_START, WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.opc.constants import RELATIONSHIP_TYPE as RT


def shade(cell, color):
    el=OxmlElement('w:shd');el.set(qn('w:fill'),color)
    cell._tc.get_or_add_tcPr().append(el)


def inline(p, text, base):
    pattern=r'(\*\*[^*]+\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))'
    for token in re.split(pattern,text):
        if not token:continue
        if token.startswith('**'):
            run=p.add_run(token[2:-2]);run.bold=True
        elif token.startswith('`'):
            run=p.add_run(token[1:-1]);run.font.name='DejaVu Sans Mono';run.font.size=Pt(8)
        elif re.fullmatch(r'\[[^\]]+\]\([^)]+\)',token):
            m=re.fullmatch(r'\[([^\]]+)\]\(([^)]+)\)',token);label,url=m.groups()
            if not re.match(r'^https?://',url):url=(base/url).resolve().as_uri()
            rel=p.part.relate_to(url,RT.HYPERLINK,is_external=True)
            link=OxmlElement('w:hyperlink');link.set(qn('r:id'),rel)
            run=OxmlElement('w:r');pr=OxmlElement('w:rPr')
            color=OxmlElement('w:color');color.set(qn('w:val'),'195C80');pr.append(color)
            run.append(pr);t=OxmlElement('w:t');t.text=label;run.append(t);link.append(run);p._p.append(link)
        else:p.add_run(token)
    if '【待补' in text:
        pr=p._p.get_or_add_pPr();el=OxmlElement('w:shd');el.set(qn('w:fill'),'FFF2CC');pr.append(el)


def configure_section(section, landscape=False):
    section.orientation=WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
    section.page_width=Cm(29.7 if landscape else 21)
    section.page_height=Cm(21 if landscape else 29.7)
    section.top_margin=Cm(1.8);section.bottom_margin=Cm(1.8)
    section.left_margin=Cm(2);section.right_margin=Cm(2)
    section.header_distance=Cm(.7);section.footer_distance=Cm(.8)


def build(source, output, pdf=True):
    doc=Document();configure_section(doc.sections[0])
    normal=doc.styles['Normal'];normal.font.name='Liberation Serif';normal.font.size=Pt(10.5)
    normal.element.rPr.rFonts.set(qn('w:eastAsia'),'Noto Serif CJK SC')
    lang=OxmlElement('w:lang');lang.set(qn('w:val'),'zh-CN');lang.set(qn('w:eastAsia'),'zh-CN');normal.element.rPr.append(lang)
    normal.paragraph_format.line_spacing=1.22;normal.paragraph_format.space_after=Pt(6)
    for name,size in [('Title',21),('Heading 1',15),('Heading 2',12),('Heading 3',11)]:
        style=doc.styles[name];style.font.name='Noto Sans CJK SC';style.font.size=Pt(size)
        style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Noto Sans CJK SC')
        style.font.color.rgb=RGBColor.from_string('153F59')
        style.paragraph_format.keep_with_next=True
    header=doc.sections[0].header.paragraphs[0];header.add_run('JIT · 中文论文工作稿 · 2026-10-11').font.size=Pt(8)
    footer=doc.sections[0].footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run('第 ')
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    footer.add_run(' 页')
    doc.core_properties.title='JIT：面向跳跃策略能力扩展的可学习探索与生成式恢复指导'
    doc.core_properties.subject='中文研究草稿；真实开发证据与待补实验分开标注'
    doc.core_properties.author='DVGC / JIT Research'
    lines=source.read_text().splitlines();i=0;landscape=False;images=[];equations=0;tables=0
    while i<len(lines):
        line=lines[i].strip();i+=1
        if not line:continue
        if line in ('<!-- landscape -->','<!-- portrait -->'):
            landscape=line=='<!-- landscape -->'
            configure_section(doc.add_section(WD_SECTION_START.NEW_PAGE),landscape);continue
        if line=='<!-- pagebreak -->':doc.add_page_break();continue
        if line.startswith('<!--'):continue
        if line.startswith('#'):
            m=re.match(r'^(#+)\s+(.*)',line);level=len(m[1])
            p=doc.add_paragraph(style='Title' if level==1 else 'Heading '+str(min(3,level-1)))
            inline(p,m[2],source.parent);continue
        if line=='$$':
            parts=[]
            while i<len(lines) and lines[i].strip()!='$$':parts.append(lines[i].strip());i+=1
            i+=1;p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            math=OxmlElement('m:oMath');r=OxmlElement('m:r')
            pr=OxmlElement('m:rPr');pr.append(OxmlElement('m:nor'));r.append(pr)
            t=OxmlElement('m:t');t.text=' '.join(parts);r.append(t);math.append(r);p._p.append(math)
            equations+=1;continue
        if line.startswith('|'):
            rows=[line]
            while i<len(lines) and lines[i].strip().startswith('|'):rows.append(lines[i].strip());i+=1
            rows=[r for r in rows if not re.fullmatch(r'[| :\-]+',r)]
            data=[[c.strip() for c in r.strip('|').split('|')] for r in rows]
            n=len(data[0]);assert all(len(r)==n for r in data),data
            table=doc.add_table(rows=0, cols=n);table.style='Table Grid';table.autofit=False
            width=Cm((25.7 if landscape else 17)/n)
            for column in table.columns:column.width=width
            for j,row in enumerate(data):
                cells=table.add_row().cells
                trpr=table.rows[-1]._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
                if j==0:trpr.append(OxmlElement('w:tblHeader'))
                for cell,txt in zip(cells,row):
                    cell.width=width;p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(3);p.paragraph_format.space_before=Pt(3)
                    p.paragraph_format.line_spacing=1.08;p.paragraph_format.keep_with_next=j<len(data)-1
                    inline(p,txt,source.parent)
                    for run in p.runs:run.font.size=Pt(9 if n<6 else 8)
                    if j==0:
                        shade(cell,'E6EEF4')
                        for run in p.runs:run.bold=True
                    elif '待补' in txt:shade(cell,'FFF2CC')
            tables+=1;doc.add_paragraph().paragraph_format.space_after=Pt(0);continue
        m=re.fullmatch(r'!\[([^\]]*)\]\(([^)]+)\)',line)
        if m:
            path=(source.parent/m[2]).resolve();assert path.is_file(),path
            p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True
            p.add_run().add_picture(str(path),width=Cm(25.5 if landscape else 16.8));images.append(str(path));continue
        if re.match(r'^[-*]\s+',line):
            p=doc.add_paragraph(style='List Bullet');inline(p,line[2:],source.parent);continue
        if re.match(r'^\d+\.\s',line):
            p=doc.add_paragraph();inline(p,line,source.parent);continue
        p=doc.add_paragraph()
        if line.startswith('> '):line=line[2:]
        inline(p,line,source.parent)
        if line.startswith(('**表','**算法')):p.paragraph_format.keep_with_next=True
        if line.startswith('**图') or line.startswith('**表'):
            for run in p.runs:run.font.size=Pt(9)
    output.mkdir(parents=True,exist_ok=True);target=output/(source.stem+'.docx');doc.save(target)
    record=dict(source=str(source),source_sha256=hashlib.sha256(source.read_bytes()).hexdigest(),
        docx=str(target),docx_sha256=hashlib.sha256(target.read_bytes()).hexdigest(),
        images=images,image_count=len(images),equation_count=equations,table_count=tables,
        simulator_interactions=0,training_updates=0,math='editable Office Math linear notation',pdf_built=False)
    if pdf:
        with tempfile.TemporaryDirectory(prefix='jit-word-lo-') as tmp:
            profile=Path(tmp)/'profile'
            lo_env=os.environ.copy()
            lo_lib=Path('/usr/lib/libreoffice/program')
            if lo_lib.is_dir():
                lo_env['LD_LIBRARY_PATH']=str(lo_lib)+(':'+lo_env['LD_LIBRARY_PATH'] if lo_env.get('LD_LIBRARY_PATH') else '')
            result=subprocess.run(['libreoffice','-env:UserInstallation='+profile.as_uri(),'--headless','--convert-to','pdf','--outdir',str(output),str(target)],capture_output=True,text=True,timeout=60,env=lo_env)
        pdfpath=target.with_suffix('.pdf')
        if result.returncode or not pdfpath.exists():raise RuntimeError(result.stdout+'\n'+result.stderr)
        record.update(pdf_built=True,pdf=str(pdfpath),pdf_sha256=hashlib.sha256(pdfpath.read_bytes()).hexdigest())
        record['pdfinfo']=subprocess.check_output(['pdfinfo',str(pdfpath)],text=True)
    (output/(source.stem+'_build.json')).write_text(json.dumps(record,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps({k:v for k,v in record.items() if k not in ('images','pdfinfo')},ensure_ascii=False))


if __name__=='__main__':
    ap=argparse.ArgumentParser();ap.add_argument('--source',type=Path,required=True);ap.add_argument('--output',type=Path,required=True);ap.add_argument('--docx-only',action='store_true')
    args=ap.parse_args();build(args.source.resolve(),args.output.resolve(),not args.docx_only)
