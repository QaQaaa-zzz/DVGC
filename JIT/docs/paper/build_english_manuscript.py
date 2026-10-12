#!/usr/bin/env python3
"""Build a journal-style English manuscript with editable Office Math.

The authoring source is Pandoc Markdown. Display equations may use ``\\tag{n}``;
these become centered editable OMML with a right-aligned equation number in
Word. The LaTeX companion retains the original mathematical source. This is
an authoring-only pipeline: it never imports the simulator or changes data.

Dependencies: Pandoc, python-docx, lxml, Pillow, Tectonic. A locally isolated
Pandoc installation can be supplied with --pandoc. No packages are installed
by this script. The Chinese manuscript builder is deliberately unaffected.
"""
from __future__ import annotations

import argparse
import copy
import hashlib
import json
import os
from pathlib import Path
import re
import shutil
import subprocess
import tempfile
from zipfile import ZipFile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from lxml import etree


NS = {
    "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
    "m": "http://schemas.openxmlformats.org/officeDocument/2006/math",
}
NUMBER_MARKER = "JITEQUATIONNUMBER"


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def finalize_latex(tex: str, citation_map: Path | None) -> str:
    """Add stable equation labels and editable standard LaTeX citations."""
    tex = re.sub(r"\\tag\{([^}]+)\}", lambda m: m.group() + "\\label{eq:" + m.group(1) + "}", tex)
    def compact_table(match):
        block = match.group()
        # Keep each of these bounded manuscript tables with its complete caption.
        # Longtable floats can otherwise interleave with figures and split a
        # caption across pages even when the entire table fits on a fresh page.
        if "Teacher-to-student conversion template" in block:
            fractions = iter(["0.440000", "0.140000", "0.140000", "0.140000", "0.140000"])
            block = re.sub(r"\\real\{[0-9.]+\}", lambda _: "\\real{" + next(fractions) + "}", block)
        elif "Main TEST results template" in block:
            fractions = iter(["0.240000", "0.152000", "0.152000", "0.152000", "0.152000", "0.152000"])
            block = re.sub(r"\\real\{[0-9.]+\}", lambda _: "\\real{" + next(fractions) + "}", block)
        else:
            values = [float(x) for x in re.findall(r"\\real\{([0-9.]+)\}", block)]
            if values:
                fractions = iter([x / sum(values) for x in values])
                block = re.sub(r"\\real\{[0-9.]+\}", lambda _: "\\real{" + f"{next(fractions):.8f}" + "}", block)
        block = block.replace("\\begin{longtable}[]", "\\begin{tabular}")
        block = block.replace("\\endhead\n", "")
        block = block.replace("\\bottomrule\\noalign{}\n\\endlastfoot\n", "")
        block = block.replace("\\end{longtable}", "\\bottomrule\n\\end{tabular}")
        return "\\par\\medskip\n\\noindent\\begin{minipage}{\\linewidth}\n" + block + "\n\\end{minipage}\n\\par\\medskip"
    tex = re.sub(r"(?s)\\textbf\{Table \d+\..*?\n\n\{\\def\\LTcaptype\{none\}.*?\\end\{longtable\}\n\}", compact_table, tex)
    if citation_map is None or not citation_map.is_file():
        return tex
    keys = json.loads(citation_map.read_text())
    marker = re.search(r"\\section\{References\}\\label\{references\}", tex)
    if not marker:
        return tex
    body, bibliography = tex[:marker.start()], tex[marker.end():]
    ref_pattern = r"(?ms)^\{\[\}(\d+)\{\]\}\s+(.*?)(?=^\{\[\}\d+\{\]\}|\\end\{document\})"
    references = re.findall(ref_pattern, bibliography)
    if set(number for number, _ in references) != set(keys):
        raise ValueError("The bibliography and numeric-to-key citation map differ")
    def cite(match):
        ids = [item.strip() for item in match.group(1).split(",")]
        return "\\cite{" + ",".join(keys[item] for item in ids) + "}"
    body = re.sub(r"\{\[\}([0-9]+(?:\s*,\s*[0-9]+)*)\{\]\}", cite, body)
    entries = ["\\bibitem{" + keys[number] + "}\n" + text.strip() for number, text in references]
    return body + "\\begin{thebibliography}{99}\n\n" + "\n\n".join(entries) + "\n\n\\end{thebibliography}\n\\end{document}\n"


def export_latex_project(source: Path, tex: Path, pdf: Path | None, destination: Path) -> dict:
    """Create a portable authoring folder with exact figure bytes."""
    destination.mkdir(parents=True, exist_ok=True)
    figures = destination / "figures"
    figures.mkdir(exist_ok=True)
    image_records = []
    def copy_image(match):
        options, name = match.group(1) or "", match.group(2)
        original = (source.parent / name).resolve()
        copied = figures / original.name
        if copied.exists() and sha256(copied) != sha256(original):
            raise ValueError("Portable figure filename collision: " + original.name)
        shutil.copy2(original, copied)
        image_records.append({"source": str(original), "path": "figures/" + copied.name, "sha256": sha256(copied)})
        return "\\includegraphics" + options + "{figures/" + copied.name + "}"
    portable = re.sub(r"\\includegraphics(\[[^\n]*\])?\{([^}\n]+)\}", copy_image, tex.read_text())
    (destination / "main.tex").write_text(portable)
    bib = source.parent / "generative_bridge_assets/references_verified.bib"
    if bib.is_file():
        shutil.copy2(bib, destination / "references_verified.bib")
    original_framework = source.parent / "generative_bridge_assets/JIT-core-method-network-icons.pdf"
    framework_record = None
    if original_framework.is_file():
        shutil.copy2(original_framework, destination / "original_framework.pdf")
        framework_record = {"path": "original_framework.pdf", "sha256": sha256(original_framework)}
    if pdf is not None and pdf.is_file():
        shutil.copy2(pdf, destination / "main.pdf")
    (destination / "README.md").write_text(
        "# JIT English manuscript — editable LaTeX project\n\n"
        "Edit `main.tex`; all six figures are in `figures/`. The original framework is preserved, "
        "and its unchanged vector PDF is supplied as `original_framework.pdf`.\n\n"
        "Compile with `tectonic main.tex`, or run `xelatex main.tex` twice in a current TeX Live installation. "
        "The prebuilt `main.pdf` is the viewing copy. The document uses A4 article layout, "
        "amsmath, unicode-math, graphicx, caption, float, booktabs, longtable and hyperref. "
        "Tectonic downloads missing standard packages on the first build.\n\n"
        "Citations use standard `\\cite{key}` commands and an editable `thebibliography` section, so no "
        "BibTeX pass is required. `references_verified.bib` supplies the verified bibliographic metadata "
        "for later conversion to a journal BibTeX style. Numbered equations have explicit tags and "
        "`eq:N` labels. Tables and equations are text, not images.\n\n"
        "Figures 4–6 and pending table entries are deliberately unfilled result templates. They are "
        "not simulated or measured findings. No training or evaluation was run to build this document.\n"
    )
    receipt = {"schema": "jit_portable_latex_v1", "main_tex_sha256": sha256(destination / "main.tex"),
               "figures": image_records, "figure_count": len(image_records),
               "original_framework": framework_record,
               "bibliography": "standard thebibliography with cite keys; verified BibTeX supplied",
               "pdf_sha256": sha256(destination / "main.pdf") if (destination / "main.pdf").is_file() else None}
    (destination / "manifest.json").write_text(json.dumps(receipt, indent=2) + "\n")
    return receipt


def run(args: list[str], **kwargs) -> subprocess.CompletedProcess:
    result = subprocess.run(args, text=True, capture_output=True, **kwargs)
    if result.returncode:
        raise RuntimeError("Command failed: " + args[0] + "\n" + result.stderr + result.stdout)
    return result


def prepare_markdown(source: str, word: bool) -> tuple[str, list[str]]:
    """Keep equations structural, never rasterize mathematical expressions."""
    numbers: list[str] = []

    def equation(match):
        body = match.group(1).strip()
        tags = re.findall(r"\\tag\{([^}]+)\}", body)
        if len(tags) > 1:
            raise ValueError("One equation number is allowed per display.")
        if tags:
            numbers.append(tags[0])
            if word:
                body = re.sub(r"\\tag\{[^}]+\}", "", body).strip()
                # Office Math supports aligned arrays, but common Linux Word
                # readers corrupt these. Separate display rows preserve every
                # symbol and remain editable in both Word and LibreOffice.
                aligned = re.fullmatch(r"\\begin\{aligned\}(.*?)\\end\{aligned\}([.,;]?)", body, re.S)
                if aligned:
                    inner = aligned.group(1)
                    rows, parts, depth, offset = [], [], 0, 0
                    tokens = re.finditer(r"\\begin\{[^}]+\}|\\end\{[^}]+\}|\\\\(?:\[[^]]*\])?|(?<!\\)&", inner)
                    for token in tokens:
                        parts.append(inner[offset:token.start()])
                        text = token.group()
                        if text.startswith("\\begin"):
                            depth += 1
                            parts.append(text)
                        elif text.startswith("\\end"):
                            depth -= 1
                            parts.append(text)
                        elif text.startswith("\\\\") and depth == 0:
                            rows.append("".join(parts).strip())
                            parts = []
                        elif text != "&" or depth > 0:
                            parts.append(text)
                        offset = token.end()
                    parts.append(inner[offset:])
                    rows.append("".join(parts).strip())
                    rows = [row for row in rows if row]
                    rows[-1] += aligned.group(2)
                    return "\n\n".join("$$\n" + row + "\n$$" for row in rows) + "\n\n" + NUMBER_MARKER + tags[0] + "\n"
                return "$$\n" + body + "\n$$\n\n" + NUMBER_MARKER + tags[0] + "\n"
        return "$$\n" + body + "\n$$"

    # Legacy {\rm word} is valid TeX but unsupported by Pandoc's math parser.
    # Normalize only that font declaration, preserving balanced nested groups.
    while True:
        match = re.search(r"\{\\rm\s+", source)
        if not match:
            break
        depth, end = 1, match.end()
        while depth and end < len(source):
            if source[end] == "{" and source[end-1] != "\\":
                depth += 1
            elif source[end] == "}" and source[end-1] != "\\":
                depth -= 1
            end += 1
        if depth:
            raise ValueError("Unbalanced legacy roman math group")
        source = source[:match.start()] + "\\mathrm{" + source[match.end():end-1] + "}" + source[end:]
    rendered = re.sub(r"(?ms)^\$\$[ \t]*\n(.*?)^\$\$[ \t]*$", equation, source)
    first = re.search(r"(?m)^# (.+)$", rendered)
    if first and not rendered.lstrip().startswith("---\n"):
        title = first.group(1)
        rendered = rendered[:first.start()] + rendered[first.end():]
        rendered = "---\ntitle: " + json.dumps(title) + "\nlang: en-US\n---\n" + rendered
    return rendered, numbers


def font_style(style, size=11, bold=False):
    style.font.name = "Times New Roman"
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor.from_string("17212B")
    for element in style._element.xpath("./w:rPr/w:rFonts"):
        for key in list(element.attrib):
            if "theme" in key.lower():
                del element.attrib[key]
    for border in style._element.xpath("./w:pPr/w:pBdr"):
        border.getparent().remove(border)


def reference_document(path: Path, pandoc: Path):
    # A Pandoc reference must retain its full style inventory (including Table,
    # Compact and CaptionedFigure), rather than using python-docx's defaults.
    with path.open("wb") as stream:
        subprocess.run([str(pandoc), "--print-default-data-file", "reference.docx"],
                       stdout=stream, check=True)
    doc = Document(path)
    section = doc.sections[0]
    section.page_width = Inches(8.2677)
    section.page_height = Inches(11.6929)
    section.top_margin = Inches(0.83)
    section.bottom_margin = Inches(0.83)
    section.left_margin = Inches(0.86)
    section.right_margin = Inches(0.86)
    section.header_distance = Inches(0.32)
    section.footer_distance = Inches(0.35)
    for name in ("Normal", "Body Text", "First Paragraph"):
        if name in doc.styles:
            font_style(doc.styles[name])
            fmt = doc.styles[name].paragraph_format
            fmt.space_before = Pt(0)
            fmt.space_after = Pt(6)
            fmt.line_spacing = 1.08
            fmt.widow_control = True
    font_style(doc.styles["Title"], 19, True)
    doc.styles["Title"].paragraph_format.space_after = Pt(12)
    for level, size in ((1, 13), (2, 11.5), (3, 11)):
        style = doc.styles[f"Heading {level}"]
        font_style(style, size, True)
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.keep_with_next = True
    for name in ("Caption", "Image Caption", "Table Caption"):
        if name in doc.styles:
            font_style(doc.styles[name], 9.5)
    doc.core_properties.author = "DVGC / JIT Research"
    doc.core_properties.subject = "English robotics research manuscript"
    doc.save(path)


def table_rule(table, tag: str, val: str, color="AAB4BE", size="5"):
    pr = table._tbl.tblPr
    borders = pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        pr.insert_element_before(borders, "w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook")
    element = OxmlElement("w:" + tag)
    element.set(qn("w:val"), val)
    element.set(qn("w:sz"), size)
    element.set(qn("w:color"), color)
    borders.append(element)


def polish_docx(path: Path, numbers: list[str]) -> dict:
    doc = Document(path)
    # Pandoc can retain theme attributes from the reference styles; concrete
    # font names must take precedence in both Word and LibreOffice.
    for name in ("Normal", "Body Text", "First Paragraph", "Compact"):
        if name in doc.styles:
            font_style(doc.styles[name], 11)
    font_style(doc.styles["Title"], 19, True)
    for level, size in ((1, 13), (2, 11.5), (3, 11)):
        font_style(doc.styles[f"Heading {level}"], size, True)
    # Preserve balanced ordinary parentheses inside a sum/product operand.
    # Pandoc may assign only the opening delimiter to m:nary/m:e; Microsoft
    # Word tolerates this, but LibreOffice's math importer drops the summand.
    for operator in doc._element.xpath(".//m:nary"):
        operand = operator.find(qn("m:e"))
        if operand is None:
            continue
        content = "".join(operand.itertext())
        if content not in ("(", "["):
            continue
        opening, closing = ("(", ")") if content == "(" else ("[", "]")
        depth = 1
        sibling = operator.getnext()
        while sibling is not None and depth:
            following = sibling.getnext()
            text = "".join(sibling.itertext())
            depth += text.count(opening) - text.count(closing)
            operand.append(sibling)
            sibling = following
        if depth:
            raise ValueError("Unbalanced delimiter in a sum/product operand")
    for script in doc._element.xpath(".//m:sSup|.//m:sSub|.//m:sSubSup"):
        base = script.find(qn("m:e"))
        if base is None:
            continue
        content = "".join(base.itertext())
        if content not in (")", "]"):
            continue
        opening, closing = ("(", ")") if content == ")" else ("[", "]")
        depth, previous, moved = 1, script.getprevious(), []
        while previous is not None and depth:
            earlier = previous.getprevious()
            text = "".join(previous.itertext())
            depth += text.count(closing) - text.count(opening)
            moved.append(previous)
            previous = earlier
        if not depth:
            for index, item in enumerate(reversed(moved)):
                base.insert(index, item)
    found_numbers = []
    for paragraph in list(doc.paragraphs):
        text = paragraph.text.strip()
        if text.startswith(NUMBER_MARKER):
            number = text[len(NUMBER_MARKER):]
            prev = paragraph._p.getprevious()
            while prev is not None and prev.tag != qn("w:p"):
                prev = prev.getprevious()
            if prev is None or not prev.xpath(".//m:oMath"):
                raise ValueError(f"Equation {number} does not follow an Office Math paragraph")
            ppr = prev.get_or_add_pPr()
            for old in list(prev):
                if old.tag == qn("m:oMathPara"):
                    index = list(prev).index(old)
                    for math in list(old):
                        if math.tag == qn("m:oMath"):
                            prev.insert(index, copy.deepcopy(math))
                            index += 1
                    prev.remove(old)
            # A centered tab and a right tab make numbering independent of math length.
            from docx.text.paragraph import Paragraph
            eq = Paragraph(prev, paragraph._parent)
            eq.alignment = WD_ALIGN_PARAGRAPH.LEFT
            eq.paragraph_format.space_before = Pt(5)
            eq.paragraph_format.space_after = Pt(8)
            eq.paragraph_format.keep_together = True
            eq.paragraph_format.tab_stops.add_tab_stop(Inches(3.27), WD_TAB_ALIGNMENT.CENTER)
            eq.paragraph_format.tab_stops.add_tab_stop(Inches(6.54), WD_TAB_ALIGNMENT.RIGHT)
            left = OxmlElement("w:r")
            left.append(OxmlElement("w:tab"))
            prev.insert(list(prev).index(ppr) + 1, left)
            eq.add_run("\t(" + number + ")").font.size = Pt(10)
            paragraph._p.getparent().remove(paragraph._p)
            found_numbers.append(number)
        elif paragraph.style.name.startswith("Heading"):
            paragraph.paragraph_format.keep_with_next = True
        elif paragraph.style.name in ("Title", "Subtitle"):
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif text.startswith(("Figure ", "Fig. ", "Table ")):
            for r in paragraph.runs:
                r.font.size = Pt(9.5)
            paragraph.paragraph_format.space_after = Pt(8)
            if text.startswith("Table "):
                paragraph.paragraph_format.keep_with_next = True
        elif text == "JITPAGEBREAK":
            paragraph.clear()
            paragraph.add_run().add_break(WD_BREAK.PAGE)
        elif paragraph._p.xpath(".//w:drawing"):
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.keep_with_next = True
            paragraph.paragraph_format.space_after = Pt(3)
        else:
            if paragraph.style.name in ("Normal", "Body Text", "First Paragraph"):
                paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if found_numbers != numbers:
        raise ValueError(f"Equation numbering changed: {found_numbers} != {numbers}")
    for table in doc.tables:
        table.autofit = False
        widths = [1 / len(table.columns)] * len(table.columns)
        if [cell.text.strip() for cell in table.rows[0].cells] == ["Quantity", "M1", "M2", "M3", "M4"]:
            widths = [.44, .14, .14, .14, .14]
        for col, weight in zip(table.columns, widths):
            col.width = Inches(6.54 * weight)
        for tag in ("left", "right", "insideV"):
            table_rule(table, tag, "nil")
        for tag in ("top", "bottom", "insideH"):
            table_rule(table, tag, "single")
        for row_index, row in enumerate(table.rows):
            pr = row._tr.get_or_add_trPr()
            pr.append(OxmlElement("w:cantSplit"))
            if row_index == 0:
                pr.append(OxmlElement("w:tblHeader"))
            for cell, weight in zip(row.cells, widths):
                cell.width = Inches(6.54 * weight)
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.keep_with_next = row_index == 0
                    paragraph.paragraph_format.space_after = Pt(3)
                    paragraph.paragraph_format.space_before = Pt(3)
                    paragraph.paragraph_format.line_spacing = 1.0
                    for r in paragraph.runs:
                        r.font.size = Pt(9)
                        if row_index == 0:
                            r.bold = True
                if row_index == 0:
                    shd = OxmlElement("w:shd")
                    shd.set(qn("w:fill"), "EEF2F5")
                    cell._tc.get_or_add_tcPr().append(shd)
    for section in doc.sections:
        head = section.header.paragraphs[0]
        head.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r = head.add_run("JIT: Generative Recovery Guidance")
        r.font.name = "Times New Roman"
        r.font.size = Pt(8)
        r.font.color.rgb = RGBColor.from_string("68717B")
        foot = section.footer.paragraphs[0]
        foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), "PAGE")
        foot._p.append(field)
    doc.save(path)
    with ZipFile(path) as archive:
        root = etree.fromstring(archive.read("word/document.xml"))
        text = " ".join(root.xpath("//w:t/text()", namespaces=NS))
        return {
            "numbered_equations": len(numbers),
            "equation_numbers": numbers,
            "editable_math_objects": len(root.xpath("//m:oMath", namespaces=NS)),
            "editable_fractions": len(root.xpath("//m:f", namespaces=NS)),
            "editable_subscripts": len(root.xpath("//m:sSub|//m:sSubSup", namespaces=NS)),
            "table_count": len(doc.tables),
            "inline_image_count": len(doc.inline_shapes),
            "leftover_number_markers": NUMBER_MARKER in text,
        }


def build(source: Path, output: Path, pandoc: Path, pdf=True, tectonic: Path | None = None, project_dir: Path | None = None):
    source = source.resolve()
    output = output.resolve()
    output.mkdir(parents=True, exist_ok=True)
    original = source.read_text()
    original_sha256 = hashlib.sha256(original.encode()).hexdigest()
    heading_shift = [] if original.lstrip().startswith("---\n") else ["--shift-heading-level-by=-1"]
    word_md, numbers = prepare_markdown(original, word=True)
    latex_md, latex_numbers = prepare_markdown(original, word=False)
    if numbers != latex_numbers:
        raise ValueError("Word and LaTeX equation inventories differ")
    target = output / (source.stem + ".docx")
    tex = output / (source.stem + ".tex")
    with tempfile.TemporaryDirectory(prefix="jit-english-manuscript-") as td:
        temp = Path(td)
        reference = temp / "reference.docx"
        reference_document(reference, pandoc)
        md = temp / "word.md"
        md.write_text(word_md.replace("<!-- pagebreak -->", "JITPAGEBREAK"))
        args = [str(pandoc), str(md), "--from=markdown+tex_math_dollars", "--standalone", *heading_shift,
                "--resource-path=" + str(source.parent), "--reference-doc=" + str(reference),
                "--output=" + str(target)]
        conversion = run(args, cwd=source.parent)
        if "Could not convert TeX math" in conversion.stderr:
            raise ValueError("Unsupported mathematical source:\n" + conversion.stderr)
        checks = polish_docx(target, numbers)
        md.write_text(latex_md.replace("<!-- pagebreak -->", "\\newpage"))
        header = temp / "journal-layout.tex"
        header.write_text(
            "\\usepackage{caption}\n"
            "\\usepackage{float}\n"
            "\\floatplacement{figure}{H}\n"
            "\\captionsetup{labelformat=empty,font=small}\n"
            "\\AtBeginDocument{\\renewcommand{\\setminus}{\\mathbin{\\backslash}}}\n"
            "\\setlength{\\floatsep}{12pt plus 2pt minus 2pt}\n"
            "\\setlength{\\textfloatsep}{14pt plus 2pt minus 2pt}\n"
        )
        latex = run([str(pandoc), str(md), "--from=markdown+tex_math_dollars", "--standalone", *heading_shift,
                     "--to=latex", "--variable=fontsize:11pt", "--variable=papersize:a4", "--variable=geometry:margin=0.86in",
                     "--include-in-header=" + str(header),
                     "--resource-path=" + str(source.parent), "--output=" + str(tex)], cwd=source.parent)
        citation_map = source.parent / "generative_bridge_assets/english/citation_keys.json"
        tex.write_text(finalize_latex(tex.read_text(), citation_map))
        if pdf:
            if tectonic is None or not tectonic.is_file():
                raise ValueError("Tectonic unavailable; provide --tectonic /path/to/tectonic")
            export = run([str(tectonic), "--keep-logs", "--outdir", str(output), str(tex)],
                         cwd=source.parent, timeout=900)
    record = {
        "schema": "jit_english_manuscript_build_v1",
        "source": str(source), "source_sha256": original_sha256,
        "source_changed_during_build": sha256(source) != original_sha256,
        "docx": str(target), "docx_sha256": sha256(target),
        "latex": str(tex), "latex_sha256": sha256(tex),
        "pandoc_version": run([str(pandoc), "--version"]).stdout.splitlines()[0],
        "docx_conversion_warnings": conversion.stderr,
        "latex_conversion_warnings": latex.stderr,
        "math": "Editable structured OMML, converted from LaTeX by Pandoc",
        "pdf_engine": str(tectonic) if pdf else None,
        "new_physical_steps": 0, "new_optimizer_updates": 0,
        "checks": checks,
    }
    if pdf:
        pdf_path = target.with_suffix(".pdf")
        if not pdf_path.exists():
            raise RuntimeError("Tectonic did not create PDF: " + export.stdout)
        record.update(pdf=str(pdf_path), pdf_sha256=sha256(pdf_path),
                      pdfinfo=run(["pdfinfo", str(pdf_path)]).stdout,
                      pdf_build_messages=export.stdout + export.stderr)
    if project_dir:
        record["portable_latex"] = export_latex_project(source, tex, target.with_suffix(".pdf") if pdf else None, project_dir.resolve())
    receipt = output / (source.stem + "_build.json")
    receipt.write_text(json.dumps(record, indent=2) + "\n")
    print(json.dumps({k: v for k, v in record.items() if k != "pdfinfo"}, indent=2))


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--pandoc", type=Path, default=Path(shutil.which("pandoc") or "/tmp/jit_manuscript_tools/pypandoc/files/pandoc"))
    parser.add_argument("--docx-only", action="store_true")
    parser.add_argument("--tectonic", type=Path, default=Path(shutil.which("tectonic") or "/tmp/jit_tectonic/tectonic"))
    parser.add_argument("--project-dir", type=Path, help="Optional portable LaTeX authoring folder")
    options = parser.parse_args()
    if not options.pandoc.is_file():
        parser.error("Pandoc unavailable; provide --pandoc /path/to/pandoc")
    build(options.source, options.output, options.pandoc, not options.docx_only, options.tectonic, options.project_dir)
